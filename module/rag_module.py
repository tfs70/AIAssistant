
import hashlib
import json
import shutil
from pathlib import Path
from typing import Final

from docx import Document as DocxDocument
from langchain_chroma import Chroma
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_core.documents import Document
from langchain_ollama import OllamaEmbeddings
from langchain_text_splitters import RecursiveCharacterTextSplitter

from module.embedding_module import EMBEDDING_MODEL_NAME


PROJECT_ROOT: Final[Path] = Path(__file__).resolve().parent.parent

DOCUMENTS_DIRECTORY: Final[Path] = PROJECT_ROOT / "documents"
PERSIST_DIRECTORY: Final[Path] = PROJECT_ROOT / "data" / "chroma_db"
MANIFEST_FILE_PATH: Final[Path] = (
    PROJECT_ROOT / "data" / "rag_index_manifest.json"
)

COLLECTION_NAME: Final[str] = "rag_documents"

SUPPORTED_EXTENSIONS: Final[set[str]] = {
    ".pdf",
    ".docx",
    ".txt",
    ".md",
}

CHUNK_SIZE: Final[int] = 900
CHUNK_OVERLAP: Final[int] = 150
BATCH_SIZE: Final[int] = 100


def create_vector_store(
    collection_name: str = COLLECTION_NAME,
    persist_directory: str | Path = PERSIST_DIRECTORY,
) -> Chroma:
    """Create or open a persistent Chroma collection."""

    embeddings = OllamaEmbeddings(
        model=EMBEDDING_MODEL_NAME,
    )

    return Chroma(
        embedding_function=embeddings,
        collection_name=collection_name,
        persist_directory=str(persist_directory),
    )


def _get_source_files() -> list[Path]:
    """Find all supported files recursively."""

    if not DOCUMENTS_DIRECTORY.is_dir():
        raise FileNotFoundError(
            f"Documents directory not found: {DOCUMENTS_DIRECTORY}"
        )

    return sorted(
        path
        for path in DOCUMENTS_DIRECTORY.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_EXTENSIONS
    )


def _relative_source(file_path: Path) -> str:
    return file_path.relative_to(
        DOCUMENTS_DIRECTORY
    ).as_posix()


def _calculate_file_hash(file_path: Path) -> str:
    """Calculate SHA-256 without loading the entire file into memory."""

    digest = hashlib.sha256()

    with file_path.open("rb") as file:
        for block in iter(lambda: file.read(1024 * 1024), b""):
            digest.update(block)

    return digest.hexdigest()


def _load_manifest() -> dict:
    if not MANIFEST_FILE_PATH.exists():
        return {}

    with MANIFEST_FILE_PATH.open(
        "r",
        encoding="utf-8",
    ) as file:
        manifest = json.load(file)

    if not isinstance(manifest, dict):
        raise ValueError("RAG manifest must contain a JSON object.")

    return manifest


def _save_manifest(manifest: dict) -> None:
    """Write the manifest atomically."""

    MANIFEST_FILE_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    temporary_path = MANIFEST_FILE_PATH.with_suffix(".tmp")

    with temporary_path.open(
        "w",
        encoding="utf-8",
    ) as file:
        json.dump(
            manifest,
            file,
            ensure_ascii=False,
            indent=2,
        )

    temporary_path.replace(MANIFEST_FILE_PATH)


def _load_file(file_path: Path) -> list[Document]:
    """Load a supported file and preserve its source metadata."""

    extension = file_path.suffix.lower()
    relative_source = _relative_source(file_path)

    if extension == ".pdf":
        documents = PyPDFLoader(str(file_path)).load()

        for document in documents:
            # PyPDFLoader normally returns a zero-based page index.
            page_index = document.metadata.get("page")

            if page_index is not None:
                document.metadata["page_number"] = int(page_index) + 1

    elif extension in {".txt", ".md"}:
        documents = TextLoader(
            str(file_path),
            encoding="utf-8",
        ).load()

    elif extension == ".docx":
        word_document = DocxDocument(str(file_path))

        parts = [
            paragraph.text.strip()
            for paragraph in word_document.paragraphs
            if paragraph.text.strip()
        ]

        # Also extract text from Word tables.
        for table in word_document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]

                if cells:
                    parts.append(" | ".join(cells))

        text = "\n".join(parts)

        documents = (
            [Document(page_content=text)]
            if text.strip()
            else []
        )

    else:
        raise ValueError(
            f"Unsupported file extension: {extension}"
        )

    for document in documents:
        document.metadata["source"] = str(file_path)
        document.metadata["relative_source"] = relative_source
        document.metadata["file_name"] = file_path.name

    return documents


def _split_documents(
    documents: list[Document],
) -> list[Document]:
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHUNK_SIZE,
        chunk_overlap=CHUNK_OVERLAP,
        separators=[
            "\n\n",
            "\n",
            "。",
            ".",
            "؛",
            "،",
            " ",
            "",
        ],
    )

    return splitter.split_documents(documents)


def _prepare_file(
    file_path: Path,
) -> tuple[str, list[Document]]:
    """Load, split and assign stable IDs to a file's chunks."""

    relative_source = _relative_source(file_path)
    documents = _load_file(file_path)

    nonempty_documents = [
        document
        for document in documents
        if document.page_content.strip()
    ]

    chunks = _split_documents(nonempty_documents)

    for index, chunk in enumerate(chunks):
        identity = (
            f"{relative_source}\n"
            f"{index}\n"
            f"{chunk.page_content}"
        )

        chunk_id = hashlib.sha256(
            identity.encode("utf-8")
        ).hexdigest()

        chunk.metadata["chunk_id"] = chunk_id

    return _calculate_file_hash(file_path), chunks


def _delete_source(
    vector_store: Chroma,
    relative_source: str,
) -> None:
    """Remove all indexed chunks belonging to one source file."""

    vector_store.delete(
        where={"relative_source": relative_source}
    )


def _add_chunks(
    vector_store: Chroma,
    chunks: list[Document],
) -> None:
    for start in range(0, len(chunks), BATCH_SIZE):
        batch = chunks[start:start + BATCH_SIZE]

        vector_store.add_documents(
            documents=batch,
            ids=[
                document.metadata["chunk_id"]
                for document in batch
            ],
        )


def rebuild_vector_store(
    clear_existing: bool = True,
) -> dict[str, int]:
    """
    Rebuild or incrementally synchronize the RAG index.

    clear_existing=True:
        Recreate the entire Chroma index from documents.

    clear_existing=False:
        Index new or changed files, skip unchanged files,
        and remove files that no longer exist.
    """

    files = _get_source_files()

    if not files:
        raise FileNotFoundError(
            "No supported documents found under "
            f"{DOCUMENTS_DIRECTORY}"
        )

    print("\n===== RAG INDEX =====")
    print(f"Directory: {DOCUMENTS_DIRECTORY}")
    print(f"Supported files found: {len(files)}")
    print(f"Embedding model: {EMBEDDING_MODEL_NAME}")
    print(f"Clear existing index: {clear_existing}")

    # Prepare all new/changed files before changing the index.
    old_manifest = (
        {} if clear_existing else _load_manifest()
    )

    prepared: dict[str, dict] = {}
    skipped = 0

    for file_path in files:
        relative_source = _relative_source(file_path)
        current_hash = _calculate_file_hash(file_path)
        previous = old_manifest.get(relative_source)

        if (
            not clear_existing
            and previous
            and previous.get("sha256") == current_hash
        ):
            skipped += 1
            print(f"[SKIP] {relative_source}")
            continue

        actual_hash, chunks = _prepare_file(file_path)

        # Detect changes while the file was being read.
        if actual_hash != current_hash:
            raise RuntimeError(
                f"File changed while being indexed: {file_path}"
            )

        if not chunks:
            raise RuntimeError(
                f"No extractable text in: {relative_source}. "
                "If this is a scanned PDF, OCR may be required."
            )

        prepared[relative_source] = {
            "sha256": actual_hash,
            "chunks": chunks,
            "chunk_ids": [
                chunk.metadata["chunk_id"]
                for chunk in chunks
            ],
        }

        print(
            f"[PREPARED] {relative_source} "
            f"| chunks={len(chunks)}"
        )

    # Do not silently continue if an incremental run has no
    # manifest but finds an old, populated collection.
    if not clear_existing and not MANIFEST_FILE_PATH.exists():
        if PERSIST_DIRECTORY.exists():
            existing_store = create_vector_store()
            if existing_store._collection.count() > 0:
                raise RuntimeError(
                    "An existing Chroma index was found without "
                    "a manifest. Run rebuild_vector_store("
                    "clear_existing=True) once."
                )

    if clear_existing:
        # Destructive operation is performed only after all files
        # have been loaded and split successfully.
        if PERSIST_DIRECTORY.exists():
            shutil.rmtree(PERSIST_DIRECTORY)

        if MANIFEST_FILE_PATH.exists():
            MANIFEST_FILE_PATH.unlink()

        vector_store = create_vector_store()
        manifest: dict = {}
        added_files = 0
        updated_files = 0

        for relative_source, item in prepared.items():
            _add_chunks(vector_store, item["chunks"])

            manifest[relative_source] = {
                "sha256": item["sha256"],
                "chunk_ids": item["chunk_ids"],
            }
            added_files += 1

    else:
        vector_store = create_vector_store()
        manifest = dict(old_manifest)
        added_files = 0
        updated_files = 0

        current_sources = {
            _relative_source(file_path)
            for file_path in files
        }

        # Remove documents that were deleted from the source folder.
        deleted_sources = (
            set(manifest.keys()) - current_sources
        )

        for relative_source in deleted_sources:
            _delete_source(vector_store, relative_source)
            manifest.pop(relative_source, None)
            print(f"[DELETED] {relative_source}")

        for relative_source, item in prepared.items():
            was_indexed = relative_source in manifest

            if was_indexed:
                _delete_source(vector_store, relative_source)
                updated_files += 1
            else:
                added_files += 1

            _add_chunks(vector_store, item["chunks"])

            manifest[relative_source] = {
                "sha256": item["sha256"],
                "chunk_ids": item["chunk_ids"],
            }

    _save_manifest(manifest)

    indexed_count = vector_store._collection.count()

    result = {
        "files_found": len(files),
        "added": added_files,
        "updated": updated_files,
        "skipped": skipped,
        "deleted": len(
            set(old_manifest.keys())
            - {_relative_source(path) for path in files}
        ) if not clear_existing else 0,
        "chunks_indexed": indexed_count,
    }

    print("\n===== INDEX RESULT =====")
    for key, value in result.items():
        print(f"{key}: {value}")

    return result


def retrieve_context(
    query: str,
    k: int = 5,
    collection_name: str = COLLECTION_NAME,
    persist_directory: str | Path = PERSIST_DIRECTORY,
) -> str:
    """Retrieve relevant chunks with source metadata."""

    if not query.strip():
        return ""

    vector_store = create_vector_store(
        collection_name=collection_name,
        persist_directory=persist_directory,
    )

    results = vector_store.similarity_search(
        query=query,
        k=k,
    )

    sections: list[str] = []

    for index, document in enumerate(results, start=1):
        source = document.metadata.get(
            "relative_source",
            document.metadata.get("source", "Unknown"),
        )
        page = document.metadata.get("page_number")

        location = f"Source: {source}"

        if page is not None:
            location += f", page: {page}"

        sections.append(
            f"[Document {index} | {location}]\n"
            f"{document.page_content}"
        )

    return "\n\n".join(sections)


def sync_rag_documents() -> dict[str, int]:
    """
    Incrementally synchronize documents with the RAG index.

    - Adds new documents.
    - Updates modified documents.
    - Skips unchanged documents.
    - Removes deleted documents.
    """
    return rebuild_vector_store(clear_existing=False)